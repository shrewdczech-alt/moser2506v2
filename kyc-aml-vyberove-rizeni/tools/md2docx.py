# -*- coding: utf-8 -*-
"""Převede dokumenty .md na .docx (pandoc) a doplní ohraničení tabulek a firemní písmo."""
import glob, os, subprocess, sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))


def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "BFBFBF")
        b.append(e)
    tblPr.append(b)
    for cell in table.rows[0].cells:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "DDEBF7")
        cell._tc.get_or_add_tcPr().append(shd)


for md in sorted(glob.glob(os.path.join(ROOT, "[0-9][0-9]_*.md"))):
    out = md[:-3] + ".docx"
    subprocess.run(["pandoc", md, "-o", out], check=True)
    d = Document(out)
    for st in d.styles:
        try:
            if st.type == 1:
                st.font.name = "Calibri"
                st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri") if st.element.rPr is not None and st.element.rPr.rFonts is not None else None
        except Exception:
            pass
    for name in ("Heading 1", "Heading 2", "Title"):
        try:
            d.styles[name].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        except KeyError:
            pass
    d.styles["Normal"].font.size = Pt(10.5)
    for t in d.tables:
        borders(t)
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Horizont Privátní banka, a.s. · Projekt SENTINEL · FIKTIVNÍ CVIČNÝ DOKUMENT"
    sec.header.paragraphs[0].runs[0].font.size = Pt(8)
    d.save(out)
    print("OK", os.path.basename(out))
