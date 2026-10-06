# -*- coding: utf-8 -*-
"""Generuje Excelové přílohy projektu SENTINEL (fiktivní výběrové řízení KYC/AML).

Spuštění:  python3 tools/generuj_xlsx.py   (z adresáře kyc-aml-vyberove-rizeni)
"""
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(__file__))
import pozadavky as P  # noqa: E402

OUT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
HESLO = "Sentinel2026"
BANKA = "Horizont Privátní banka, a.s."
PROJEKT = "Projekt SENTINEL – výměna KYC/AML řešení"
FIKCE = "FIKTIVNÍ CVIČNÝ DOKUMENT – banka, osoby a dodavatelé jsou smyšlení."

# ---------------------------------------------------------------- styly
NAVY = "1F3864"
F_HEAD = PatternFill("solid", fgColor=NAVY)
F_BANK = PatternFill("solid", fgColor="DDEBF7")     # vyplňuje Banka
F_VEND = PatternFill("solid", fgColor="FFF2CC")     # vyplňuje dodavatel
F_COMM = PatternFill("solid", fgColor="E2EFDA")     # vyplňuje komise
F_CALC = PatternFill("solid", fgColor="EDEDED")     # vzorec
F_KO = PatternFill("solid", fgColor="F8CBAD")
F_RED = PatternFill("solid", fgColor="FFC7CE")
F_GREEN = PatternFill("solid", fgColor="C6EFCE")
F_AMBER = PatternFill("solid", fgColor="FFEB9C")
WHITE_B = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14, color=NAVY)
SUB = Font(italic=True, size=9, color="7F7F7F")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
UNLOCK = Protection(locked=False)
CZK = '#,##0 "Kč"'
PCT = "0.0%"


def hlavicka(ws, row, headers, fills=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = WHITE_B
        c.fill = F_HEAD if not fills else fills[i - 1]
        if fills and fills[i - 1] is not F_HEAD:
            c.font = BOLD
        c.alignment = CENTER
        c.border = BOX


def sirky(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def titul(ws, nazev, podtitul=None):
    ws["A1"] = nazev
    ws["A1"].font = TITLE
    ws["A2"] = f"{BANKA} · {PROJEKT}"
    ws["A2"].font = Font(size=10, color="595959")
    ws["A3"] = podtitul or FIKCE
    ws["A3"].font = SUB


def navod(wb, nazev, radky):
    ws = wb.active
    ws.title = "Návod"
    titul(ws, nazev)
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 120
    r = 5
    for t in radky:
        if t.startswith("# "):
            ws.cell(row=r, column=2, value=t[2:]).font = Font(bold=True, size=12, color=NAVY)
        else:
            c = ws.cell(row=r, column=2, value=t)
            c.alignment = WRAP
        r += 1
    ws.sheet_view.showGridLines = False
    return ws


def legenda(ws, row, col=2):
    for i, (fill, txt) in enumerate([
        (F_BANK, "Vyplňuje Banka (stakeholdeři / projektový tým)"),
        (F_VEND, "Vyplňuje dodavatel"),
        (F_COMM, "Vyplňuje hodnoticí komise"),
        (F_CALC, "Vypočteno vzorcem – needitovat"),
    ]):
        ws.cell(row=row + i, column=col - 1).fill = fill
        ws.cell(row=row + i, column=col, value=txt)


def dv_list(ws, values, rng, prompt=None):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    dv.error = "Vyberte hodnotu ze seznamu."
    dv.errorTitle = "Neplatná hodnota"
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


# ---------------------------------------------------------------- Matice shody
ODPOVEDI = ["OOTB", "Konfigurace", "Custom vývoj", "3rd party", "Nesplněno"]
ODPOVEDI_LEG = ["Akceptuje", "Akceptuje s výhradou", "Neakceptuje"]
BODY = {"OOTB": 5, "Konfigurace": 3, "Custom vývoj": 1, "3rd party": 0, "Nesplněno": 0}
MOSCOW = ["Must", "Should", "Could", "Won't"]
HODNOCENO = ["K.O. (Pass/Fail)", "Byznys (30 %)", "Podklad pro FS (20 %)", "Smluvní podmínka"]
STAKEHOLDERS = ["IT", "Compliance", "Legal", "Obchod"]
STAVY = ["Návrh", "K diskusi", "Schváleno s úpravou", "Schváleno", "Zamítnuto"]

# klíč, záhlaví, šířka, skupina, varianty (I = interní, V = verze pro dodavatele)
SLOUPCE = [
    ("id", "ID požadavku", 11, "bank", "IV"),
    ("obl", "Oblast", 13, "bank", "IV"),
    ("txt", "Text požadavku", 60, "bank", "IV"),
    ("zdroj", "Zdroj / zdůvodnění (legislativa, standard)", 30, "bank", "IV"),
    ("own", "Vlastník požadavku (stakeholder)", 13, "bank", "I"),
    ("mos", "Priorita (MoSCoW)", 10, "bank", "IV"),
    ("ko", "K.O. kritérium", 9, "bank", "IV"),
    ("acc", "Akceptační kritérium (jak ověříme)", 40, "bank", "IV"),
    ("fsx", "Očekávaná kapitola FS", 10, "bank", "IV"),
    ("hod", "Hodnoceno v", 16, "bank", "IV"),
    ("stav", "Stav projednání", 13, "bank", "I"),
    ("pozn", "Poznámka z workshopu", 32, "bank", "I"),
    ("odp", "Odpověď dodavatele", 15, "vend", "IV"),
    ("kom", "Komentář dodavatele (jak je splněno, verze produktu, omezení)", 40, "vend", "IV"),
    ("fs", "Odkaz do Feasibility Study (kapitola)", 13, "vend", "IV"),
    ("chk", "Křížová kontrola FS", 14, "calc", "IV"),
    ("vaha", "Váha MoSCoW", 8, "calc", "I"),
    ("body", "Body dle odpovědi (0–5)", 9, "calc", "I"),
    ("ovr", "Ověřeno komisí (FS / demo)", 11, "comm", "I"),
    ("uzn", "Uznané body", 9, "calc", "I"),
    ("ko_v", "Výsledek K.O.", 12, "calc", "I"),
    ("wb", "Vážené body", 9, "calc", "I"),
    ("wmax", "Max. vážené body", 9, "calc", "I"),
    ("kk", "Komentář komise (povinný při snížení bodů)", 36, "comm", "I"),
]
FILL_GROUP = {"bank": F_BANK, "vend": F_VEND, "comm": F_COMM, "calc": F_CALC}


def hodnoceno_v(req):
    if req[6] == "Ano":
        return "K.O. (Pass/Fail)"
    if req[1] in P.BYZNYS:
        return "Byznys (30 %)"
    if req[1] in P.SMLUVNI or req[0] == "REG-001":
        return "Smluvní podmínka"
    return "Podklad pro FS (20 %)"


def je_pravni(req_id, oblast):
    return oblast == "LEG" or req_id == "REG-001"


def matice(path, rows, varianta, nazev, prazdnych=0, s_workshopy=True):
    wb = Workbook()
    interni = varianta == "I"
    if interni:
        nav = [
            "# Účel",
            "Matice shody je jediný zdroj požadavků výběrového řízení. Každý požadavek má unikátní ID, vlastníka, prioritu MoSCoW, "
            "zdroj (legislativa / standard Banky) a akceptační kritérium. Dodavatel ke každému požadavku vybírá normovanou odpověď, "
            "takže odpovědi lze kvantitativně porovnat.",
            "# Jak sbírat požadavky se stakeholdery",
            "1) Business analytik připraví návrh požadavků z analýzy výchozího stavu (procesy onboardingu, screeningu, rozhodování v CRM).",
            "2) Samostatné workshopy po stakeholderech: IT, Compliance, Legal, Obchod. Každý workshop zapsat do listu Workshopy.",
            "3) Každý požadavek má jednoho vlastníka, který ho obhájí a schválí jeho akceptační kritérium.",
            "4) Konflikty mezi stakeholdery (např. Obchod vs. Compliance) zapsat do listu Konflikty a rozhodnout na konsolidačním workshopu / SteerCo.",
            "5) Před vyhlášením tendru musí mít všechny požadavky stav Schváleno nebo Schváleno s úpravou; K.O. kritéria a váhy schvaluje SteerCo.",
            "6) Pro dodavatele se z matice odstraní interní sloupce (vlastník, stav, poznámka, hodnoticí sloupce) – viz Příloha 1 RFP.",
            "# Pravidla psaní požadavků",
            "• Jeden požadavek = jedna ověřitelná vlastnost. Žádné „systém bude uživatelsky přívětivý“.",
            "• Akceptační kritérium říká, jak splnění ověříme (demo, test v SIT, doklad).",
            "• ID má tvar OBLAST-NNN (KYC, SCR, TM, CAS, REP, INT, MIG, NFR, SEC, REG, LEG); ID se po vyhlášení nemění, zrušený požadavek dostane stav Zamítnuto.",
            "• K.O. = Pass/Fail, nebodují se. Must = 3, Should = 2, Could = 1, Won't = 0 (váha v bodování).",
            "# Normované odpovědi dodavatele a body",
            "OOTB = 5 b. (funkce je v základu, stačí zapnout) · Konfigurace = 3 b. (nastavení v UI bez programování) · "
            "Custom vývoj = 1 b. (nutný vývoj, zvyšuje technický dluh a TCO) · 3rd party = 0 b. · Nesplněno = 0 b.",
            "Smluvní podmínky (LEG, REG-001): Akceptuje / Akceptuje s výhradou / Neakceptuje.",
            "# Pravidlo křížové kontroly",
            "U odpovědí OOTB, Konfigurace a Custom vývoj musí dodavatel uvést kapitolu Feasibility Study, kde je řešení popsáno. "
            "Pokud odkaz chybí, nebo komise při kontrole FS / skriptovaném demu tvrzení neověří (sloupec Ověřeno = Ne), "
            "požadavek se považuje za nesplněný (0 bodů; u K.O. = FAIL). Každé snížení bodů komise zdůvodní jednou větou.",
            "# Barvy",
        ]
        ws = navod(wb, nazev, nav)
        legenda(ws, 5 + len(nav) + 1)
    else:
        nav = [
            "# Pokyny pro uchazeče",
            "1) Ke každému požadavku vyberte ve sloupci „Odpověď dodavatele“ právě jednu normovanou hodnotu.",
            "   OOTB = funkce je ve standardním produktu, stačí zapnout · Konfigurace = nastavení v UI bez programování · "
            "Custom vývoj = vyžaduje vývoj nad rámec produktu · 3rd party = zajišťuje produkt třetí strany · Nesplněno.",
            "   U smluvních podmínek (LEG-*, REG-001): Akceptuje / Akceptuje s výhradou / Neakceptuje. Výhradu popište v komentáři.",
            "2) Ve sloupci „Komentář dodavatele“ stručně popište, jak je požadavek splněn (modul, verze, omezení).",
            "3) U odpovědí OOTB, Konfigurace a Custom vývoj je POVINNÝ odkaz na kapitolu Feasibility Study (Příloha 2). "
            "Bez odkazu se požadavek považuje za nesplněný. Sloupec „Křížová kontrola FS“ vám chybějící odkazy ukáže.",
            "4) Uvedená tvrzení budou ověřena ve skriptovaném demu a jsou závazná – stanou se součástí smlouvy.",
            "5) Nevyplňujte ani neměňte modré buňky. List je uzamčen; editovatelné jsou pouze žluté buňky.",
            "6) K.O. kritéria se hodnotí Pass/Fail. Nesplnění kteréhokoli K.O. kritéria znamená vyřazení nabídky.",
            "7) Dotazy k matici zasílejte výhradně kontaktní osobě dle kap. 8 Zadávací dokumentace; odpovědi jsou sdíleny všem uchazečům (Q&A log).",
            "# Barvy",
        ]
        ws = navod(wb, nazev, nav)
        legenda(ws, 5 + len(nav) + 1)

    cols = [c for c in SLOUPCE if varianta in c[4]]
    L = {c[0]: get_column_letter(i) for i, c in enumerate(cols, 1)}
    ws = wb.create_sheet("Požadavky")
    titul(ws, "Matice shody – požadavky" + ("" if interni else " (Příloha 1 Zadávací dokumentace)"))
    HR = 5
    hlavicka(ws, HR, [c[1] for c in cols], [FILL_GROUP[c[3]] if c[3] != "bank" else F_HEAD for c in cols])
    # skupinové záhlaví
    for i, c in enumerate(cols, 1):
        g = {"bank": "BANKA", "vend": "DODAVATEL", "comm": "KOMISE", "calc": "VÝPOČET"}[c[3]]
        cell = ws.cell(row=HR - 1, column=i, value=g)
        cell.fill = FILL_GROUP[c[3]]
        cell.font = BOLD
        cell.alignment = CENTER
    ws.row_dimensions[HR].height = 48
    sirky(ws, [c[2] for c in cols])

    data = list(rows) + [None] * prazdnych
    first = HR + 1
    last = HR + len(data)
    for n, req in enumerate(data):
        r = first + n
        vals = {}
        if req:
            vals = dict(id=req[0], obl=req[1], txt=req[2], zdroj=req[3], own=req[4], mos=req[5], ko=req[6],
                        acc=req[7], fsx=req[8], hod=hodnoceno_v(req), stav=req[9], pozn=req[10])
        def ref(k):
            return f"{L[k]}{r}"
        fx = {
            "chk": f'=IF({ref("odp")}="","",IF(OR({ref("odp")}="OOTB",{ref("odp")}="Konfigurace",{ref("odp")}="Custom vývoj"),'
                   f'IF({ref("fs")}="","CHYBÍ odkaz FS","OK"),"N/A"))',
        }
        if interni:
            fx.update({
                "vaha": f'=IF({ref("mos")}="Must",3,IF({ref("mos")}="Should",2,IF({ref("mos")}="Could",1,0)))',
                "body": f'=IF({ref("odp")}="","",IF({ref("odp")}="OOTB",5,IF({ref("odp")}="Konfigurace",3,'
                        f'IF({ref("odp")}="Custom vývoj",1,IF(OR({ref("odp")}="3rd party",{ref("odp")}="Nesplněno"),0,"")))))',
                "uzn": f'=IF({ref("body")}="","",IF(OR({ref("chk")}="CHYBÍ odkaz FS",{ref("ovr")}="Ne"),0,{ref("body")}))',
                "ko_v": f'=IF({ref("ko")}<>"Ano","",IF({ref("odp")}="","",IF(OR({ref("odp")}="Nesplněno",{ref("odp")}="3rd party",'
                        f'{ref("odp")}="Neakceptuje",{ref("ovr")}="Ne",{ref("chk")}="CHYBÍ odkaz FS"),"FAIL",'
                        f'IF({ref("odp")}="Akceptuje s výhradou","K VYJASNĚNÍ","PASS"))))',
                "wb": f'=IF(AND({ref("hod")}="Byznys (30 %)",{ref("ko")}<>"Ano",ISNUMBER({ref("uzn")})),{ref("uzn")}*{ref("vaha")},0)',
                "wmax": f'=IF(AND({ref("hod")}="Byznys (30 %)",{ref("ko")}<>"Ano",{ref("mos")}<>"Won\'t"),5*{ref("vaha")},0)',
            })
        for i, c in enumerate(cols, 1):
            k = c[0]
            cell = ws.cell(row=r, column=i)
            if k in fx:
                cell.value = fx[k]
            elif k in vals:
                cell.value = vals[k]
            cell.alignment = CENTER if c[2] <= 16 else WRAP
            cell.border = BOX
            if c[3] == "vend":
                cell.fill = F_VEND
                cell.protection = UNLOCK
            elif c[3] == "comm":
                cell.fill = F_COMM
                cell.protection = UNLOCK
            elif c[3] == "calc":
                cell.fill = F_CALC
            elif interni:
                cell.protection = UNLOCK
            if k == "ko" and vals.get("ko") == "Ano":
                cell.fill = F_KO
                cell.font = BOLD
        # validace odpovědi podle typu požadavku
        if req and je_pravni(req[0], req[1]):
            dv_list(ws, ODPOVEDI_LEG, f"{L['odp']}{r}")
        else:
            dv_list(ws, ODPOVEDI if req else ODPOVEDI + ODPOVEDI_LEG, f"{L['odp']}{r}")

    rng = lambda k: f"{L[k]}{first}:{L[k]}{last}"  # noqa: E731
    if interni:
        dv_list(ws, MOSCOW, rng("mos"))
        dv_list(ws, ["Ano", "Ne"], rng("ko"))
        dv_list(ws, P_OBL := list(P.OBLASTI.keys()), rng("obl"))
        dv_list(ws, STAKEHOLDERS, rng("own"))
        dv_list(ws, HODNOCENO, rng("hod"))
        dv_list(ws, STAVY, rng("stav"))
        dv_list(ws, ["Ano", "Ne"], rng("ovr"), "Ano = tvrzení ověřeno ve FS / demu; Ne = neověřeno → 0 bodů")
        ws.conditional_formatting.add(rng("ko_v"), CellIsRule(operator="equal", formula=['"FAIL"'], fill=F_RED))
        ws.conditional_formatting.add(rng("ko_v"), CellIsRule(operator="equal", formula=['"PASS"'], fill=F_GREEN))
        ws.conditional_formatting.add(rng("ko_v"), CellIsRule(operator="equal", formula=['"K VYJASNĚNÍ"'], fill=F_AMBER))
        ws.conditional_formatting.add(rng("stav"), CellIsRule(operator="equal", formula=['"K diskusi"'], fill=F_AMBER))
        ws.conditional_formatting.add(
            rng("kk"), FormulaRule(formula=[f'AND(ISNUMBER({L["uzn"]}{first}),{L["uzn"]}{first}<{L["body"]}{first},{L["kk"]}{first}="")'], fill=F_RED))
    ws.conditional_formatting.add(rng("chk"), CellIsRule(operator="equal", formula=['"CHYBÍ odkaz FS"'], fill=F_RED))
    ws.conditional_formatting.add(rng("chk"), CellIsRule(operator="equal", formula=['"OK"'], fill=F_GREEN))
    ws.freeze_panes = f"{L['txt']}{first}"
    ws.auto_filter.ref = f"A{HR}:{get_column_letter(len(cols))}{last}"
    ws.protection.sheet = True
    ws.protection.password = HESLO
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    if interni:
        # interní verzi zamykáme jen kvůli vzorcům; Banka edituje modré buňky
        pass

    # Souhrn
    s = wb.create_sheet("Souhrn")
    titul(s, "Souhrn matice shody")
    hlavicka(s, 5, ["Oblast", "Název oblasti", "Počet požadavků", "Must", "Should", "Could", "K.O."] +
             (["Vyplněno dodavatelem", "Custom vývoj", "Nesplněno / 3rd party", "Chybí odkaz FS"]))
    sirky(s, [10, 42, 12, 8, 8, 8, 8, 14, 12, 14, 14])
    req_rng = f"'Požadavky'!${L['obl']}${first}:${L['obl']}${last}"
    def cr(k):
        return f"'Požadavky'!${L[k]}${first}:${L[k]}${last}"
    for i, (k, nm) in enumerate(P.OBLASTI.items()):
        r = 6 + i
        s.cell(row=r, column=1, value=k)
        s.cell(row=r, column=2, value=nm)
        s.cell(row=r, column=3, value=f'=COUNTIF({req_rng},A{r})')
        for j, m in enumerate(["Must", "Should", "Could"]):
            s.cell(row=r, column=4 + j, value=f'=COUNTIFS({req_rng},A{r},{cr("mos")},"{m}")')
        s.cell(row=r, column=7, value=f'=COUNTIFS({req_rng},A{r},{cr("ko")},"Ano")')
        s.cell(row=r, column=8, value=f'=COUNTIFS({req_rng},A{r},{cr("odp")},"<>")')
        s.cell(row=r, column=9, value=f'=COUNTIFS({req_rng},A{r},{cr("odp")},"Custom vývoj")')
        s.cell(row=r, column=10, value=f'=COUNTIFS({req_rng},A{r},{cr("odp")},"Nesplněno")+COUNTIFS({req_rng},A{r},{cr("odp")},"3rd party")')
        s.cell(row=r, column=11, value=f'=COUNTIFS({req_rng},A{r},{cr("chk")},"CHYBÍ odkaz FS")')
    tr = 6 + len(P.OBLASTI)
    s.cell(row=tr, column=1, value="Celkem").font = BOLD
    for col in range(3, 12):
        cl = get_column_letter(col)
        s.cell(row=tr, column=col, value=f"=SUM({cl}6:{cl}{tr - 1})").font = BOLD
    for row in s.iter_rows(min_row=6, max_row=tr, max_col=11):
        for c in row:
            c.border = BOX
    if interni:
        k = tr + 2
        s.cell(row=k, column=1, value="Vyhodnocení (po doplnění odpovědí dodavatele)").font = Font(bold=True, color=NAVY, size=12)
        items = [
            ("Pokrytí byznys funkcionalit (vážené)", f"=IFERROR(SUM({cr('wb')})/SUM({cr('wmax')}),0)", PCT),
            ("Body za Byznys funkcionality (max. 30)", f"=B{k + 1}*30", "0.00"),
            ("Počet K.O. = FAIL", f'=COUNTIF({cr("ko_v")},"FAIL")', "0"),
            ("Počet K.O. = K VYJASNĚNÍ", f'=COUNTIF({cr("ko_v")},"K VYJASNĚNÍ")', "0"),
            ("Výsledek K.O. brány", f'=IF(COUNTIF({cr("odp")},"<>")=0,"čeká na nabídku",IF(B{k + 3}>0,"VYŘAZEN","POSTUPUJE"))', "@"),
            ("Požadavky ve stavu K diskusi (musí být 0 před vyhlášením)", f'=COUNTIF({cr("stav")},"K diskusi")+COUNTIF({cr("stav")},"Návrh")', "0"),
        ]
        for i, (lbl, f, fmt) in enumerate(items, 1):
            s.cell(row=k + i, column=1, value=lbl)
            c = s.cell(row=k + i, column=2, value=f)
            c.number_format = fmt
            c.fill = F_CALC
            c.font = BOLD
        s.column_dimensions["A"].width = 48
        s.conditional_formatting.add(f"B{k + 5}", CellIsRule(operator="equal", formula=['"VYŘAZEN"'], fill=F_RED))
        s.conditional_formatting.add(f"B{k + 6}", CellIsRule(operator="greaterThan", formula=["0"], fill=F_AMBER))

    # Číselníky
    c = wb.create_sheet("Číselníky")
    titul(c, "Číselníky")
    hlavicka(c, 5, ["Oblast", "Název", "", "Odpověď", "Body", "Význam", "", "MoSCoW", "Váha"])
    for i, (k2, v) in enumerate(P.OBLASTI.items()):
        c.cell(row=6 + i, column=1, value=k2)
        c.cell(row=6 + i, column=2, value=v)
    vyz = {"OOTB": "Funkce je ve standardu produktu, stačí zapnout",
           "Konfigurace": "Nastavení v UI bez programování",
           "Custom vývoj": "Vyžaduje vývoj; zvyšuje technický dluh a TCO",
           "3rd party": "Zajišťuje produkt třetí strany",
           "Nesplněno": "Dodavatel neumí"}
    for i, o in enumerate(ODPOVEDI):
        c.cell(row=6 + i, column=4, value=o)
        c.cell(row=6 + i, column=5, value=BODY[o])
        c.cell(row=6 + i, column=6, value=vyz[o])
    for i, (m, v) in enumerate(zip(MOSCOW, [3, 2, 1, 0])):
        c.cell(row=6 + i, column=8, value=m)
        c.cell(row=6 + i, column=9, value=v)
    sirky(c, [10, 42, 3, 16, 7, 46, 3, 10, 7])

    if interni and s_workshopy:
        w = wb.create_sheet("Workshopy")
        titul(w, "Záznam workshopů se stakeholdery")
        hlavicka(w, 5, ["ID", "Datum", "Stakeholder", "Účastníci (role)", "Témata", "Závěry", "Dotčené požadavky"])
        sirky(w, [8, 12, 13, 46, 34, 52, 24])
        src = P.WORKSHOPY if len(rows) > 3 else []
        for i in range(max(len(src), 12)):
            for j in range(7):
                cell = w.cell(row=6 + i, column=j + 1, value=src[i][j] if i < len(src) else None)
                cell.alignment = WRAP
                cell.border = BOX
        dv_list(w, STAKEHOLDERS + ["Všichni", "SteerCo"], "C6:C60")

        k3 = wb.create_sheet("Konflikty")
        titul(k3, "Konflikty požadavků mezi stakeholdery a jejich rozhodnutí")
        hlavicka(k3, 5, ["ID", "Požadavek", "Pozice A", "Pozice B", "Rozhodnutí", "Rozhodnuto na"])
        sirky(k3, [7, 11, 46, 46, 52, 13])
        src = P.KONFLIKTY if len(rows) > 3 else []
        for i in range(max(len(src), 8)):
            for j in range(6):
                cell = k3.cell(row=6 + i, column=j + 1, value=src[i][j] if i < len(src) else None)
                cell.alignment = WRAP
                cell.border = BOX

        z = wb.create_sheet("Změnový log")
        titul(z, "Změnový log matice")
        hlavicka(z, 5, ["Verze", "Datum", "Změna", "Dotčené ID", "Autor", "Schválil"])
        sirky(z, [8, 12, 60, 20, 22, 22])
        log = [("0.1", "2026-09-04", "Návrh požadavků z analýzy výchozího stavu", "—", "Business analytik", "—"),
               ("0.5", "2026-09-17", "Doplněno po workshopech WS-01 až WS-04", "—", "Business analytik", "—"),
               ("0.9", "2026-09-21", "Konsolidace, MoSCoW, K.O. kritéria; konflikty K-01 až K-03", "CAS-004, KYC-007, NFR-006", "Business analytik", "Konsolidační workshop"),
               ("1.0", "2026-09-23", "Schválení matice a váhového modelu", "—", "Business analytik", "SteerCo")] if len(rows) > 3 else []
        for i in range(max(len(log), 10)):
            for j in range(6):
                cell = z.cell(row=6 + i, column=j + 1, value=log[i][j] if i < len(log) else None)
                cell.alignment = WRAP
                cell.border = BOX

    wb.active = 1
    wb.save(path)


# ---------------------------------------------------------------- TCO
ROKY = ["Rok 1", "Rok 2", "Rok 3", "Rok 4", "Rok 5"]


def tco(path):
    wb = Workbook()
    navod(wb, "Příloha 3 – Závazná cenová šablona (TCO na 5 let)", [
        "# Pokyny",
        "1) Vyplňujte VÝHRADNĚ žlutě podbarvené buňky. Ostatní buňky jsou uzamčeny a obsahují vzorce.",
        "2) Všechny ceny uvádějte v Kč bez DPH. Ceny v EUR přepočtěte kurzem uvedeným na listu Předpoklady.",
        "3) Každý řádek musí být vyplněn. Pokud je položka zahrnuta v jiné položce, uveďte 0 a do sloupce Poznámka napište, ve které.",
        "4) Rok 1 začíná podpisem smlouvy (plánováno 1. 4. 2027). Go-live je předpokládán v roce 1 nebo 2 dle vašeho plánu ve FS.",
        "5) Ceny musí odpovídat objemům a prostředím z listu Předpoklady, včetně růstového scénáře +20 % objemů od roku 4.",
        "6) Legislativní údržba (AML zákon, vyhlášky ČNB, AMLR) a oprava zranitelností v dodaném kódu musí být zahrnuty v ceně "
        "licence/podpory – na listu C to potvrďte hodnotou „Ano“. Hodnota „Ne“ je vyřazovací.",
        "7) Rate card (list D) je závazná po celou dobu smlouvy; meziroční navýšení sazeb je omezeno inflační doložkou s maximem 5 %.",
        "8) Rozvojová kapacita (list C) je oceněna automaticky z rate card na 120 MD ročně, aby byly nabídky srovnatelné.",
        "9) Hodnocenou cenou je „TCO celkem za 5 let“ na listu Souhrn. Kontrolní buňky na listu Souhrn musí být zelené.",
        "# Proč se ptáme na tyto položky",
        "Šablona vědomě obsahuje položky, které se v nabídkách často „zapomínají“ a doplácejí se později jako vícepráce: "
        "neprodukční prostředí, licence třetích stran, datové zdroje sankcí/PEP, dopad růstu objemů, integrace na logování, "
        "monitoring a zálohování, datová migrace, dokumentace v češtině, oprava bezpečnostních nálezů a exit.",
        "# Barvy",
    ])
    legenda(wb.active, 20)

    # Předpoklady (Banka)
    p = wb.create_sheet("Předpoklady")
    titul(p, "Předpoklady pro nacenění (stanovuje Banka – needitovat)")
    hlavicka(p, 5, ["Parametr", "Hodnota", "Jednotka", "Poznámka"])
    sirky(p, [52, 16, 14, 60])
    pred = [
        ("Počet aktivních klientů (rok 1)", 9500, "klientů", "65 % FO, 35 % PO a struktury"),
        ("Počet prověřovaných subjektů (klienti, zástupci, UBO)", 38000, "subjektů", "Noční re-screening"),
        ("Počet transakcí ročně (rok 1)", 1400000, "transakcí", "Vstup do monitoringu transakcí"),
        ("Počet plateb ke screeningu ročně", 220000, "plateb", "Zahraniční platby SWIFT/SEPA"),
        ("Pojmenovaní uživatelé", 150, "uživatelů", "Obchod 100, Compliance 30, podpora 20"),
        ("Souběžní uživatelé", 60, "uživatelů", ""),
        ("Růstový scénář – nárůst objemů od roku 4", 0.20, "%", "Oceňte dopad do licencí a infrastruktury (list A, řádek A.10)"),
        ("Prostředí", "DEV, SIT, UAT, PreProd, PROD, DR", "", "Všechna musí být oceněna"),
        ("Hosting", "On-premise DC Banky / privátní cloud EU", "", "Viz NFR-006 (K.O.)"),
        ("Kurz pro přepočet EUR/CZK", 25.00, "Kč/EUR", "Pevný kurz pro účely hodnocení"),
        ("Diskontní sazba pro NPV (informativně)", 0.05, "%", "NPV se nehodnotí, slouží Bance pro interní business case"),
        ("Rozvojová kapacita ročně", 120, "MD/rok", "Oceněno z rate card (list D)"),
    ]
    for i, (a, b, c_, d) in enumerate(pred):
        r = 6 + i
        for j, v in enumerate([a, b, c_, d]):
            cell = p.cell(row=r, column=j + 1, value=v)
            cell.border = BOX
            cell.fill = F_BANK
            cell.alignment = WRAP
        if c_ == "%":
            p.cell(row=r, column=2).number_format = "0%"
        elif isinstance(b, (int, float)):
            p.cell(row=r, column=2).number_format = "#,##0.00" if isinstance(b, float) else "#,##0"
    p.protection.sheet = True
    p.protection.password = HESLO
    KURZ = "'Předpoklady'!$B$15"
    DISK = "'Předpoklady'!$B$16"
    KAPACITA = "'Předpoklady'!$B$17"

    def rocni_list(nazev, titulek, polozky, extra_col=None):
        ws = wb.create_sheet(nazev)
        titul(ws, titulek)
        heads = ["Č.", "Položka", "Popis / co musí obsahovat"] + ([extra_col[0]] if extra_col else []) + ROKY + ["Celkem 5 let", "Poznámka dodavatele"]
        hlavicka(ws, 5, heads)
        sirky(ws, [6, 38, 52] + ([16] if extra_col else []) + [14] * 5 + [16, 36])
        off = 4 if extra_col else 3
        for i, (cid, nm, desc) in enumerate(polozky):
            r = 6 + i
            ws.cell(row=r, column=1, value=cid)
            ws.cell(row=r, column=2, value=nm).alignment = WRAP
            ws.cell(row=r, column=3, value=desc).alignment = WRAP
            if extra_col:
                c = ws.cell(row=r, column=4)
                c.fill = F_VEND
                c.protection = UNLOCK
                dv_list(ws, extra_col[1], f"D{r}")
            for y in range(5):
                c = ws.cell(row=r, column=off + 1 + y)
                c.fill = F_VEND
                c.protection = UNLOCK
                c.number_format = CZK
            a, b = get_column_letter(off + 1), get_column_letter(off + 5)
            t = ws.cell(row=r, column=off + 6, value=f"=SUM({a}{r}:{b}{r})")
            t.fill = F_CALC
            t.number_format = CZK
            n = ws.cell(row=r, column=off + 7)
            n.fill = F_VEND
            n.protection = UNLOCK
            n.alignment = WRAP
            for col in range(1, off + 8):
                ws.cell(row=r, column=col).border = BOX
        tr = 6 + len(polozky)
        ws.cell(row=tr, column=2, value="Mezisoučet").font = BOLD
        for col in range(off + 1, off + 7):
            cl = get_column_letter(col)
            c = ws.cell(row=tr, column=col, value=f"=SUM({cl}6:{cl}{tr - 1})")
            c.font = BOLD
            c.fill = F_CALC
            c.number_format = CZK
            c.border = BOX
        ws.freeze_panes = "D6"
        ws.protection.sheet = True
        ws.protection.password = HESLO
        return ws, tr, off

    A = [
        ("A.1", "Licence – produkční prostředí (PROD)", "Všechny moduly pokrývající Matici shody; metrika licencování uvedena v poznámce"),
        ("A.2", "Licence – DR prostředí", "Záložní DC (active-passive)"),
        ("A.3", "Licence – DEV", "Vývojové prostředí Banky"),
        ("A.4", "Licence – SIT", "Systémové integrační testy"),
        ("A.5", "Licence – UAT", "Uživatelské akceptační testy"),
        ("A.6", "Licence – PreProd (výkonnostní)", "Prostředí v produkční velikosti pro výkonnostní testy"),
        ("A.7", "Licence databáze třetí strany", "Oracle / MS SQL / PostgreSQL – uveďte OEM v ceně nebo nákup Bankou"),
        ("A.8", "Licence aplikačního serveru, middleware, ostatní SW třetích stran", "Vše, co řešení potřebuje k běhu"),
        ("A.9", "Datové zdroje: sankční seznamy, PEP, adverse media", "Předplatné datových feedů po celých 5 let"),
        ("A.10", "Dopad růstového scénáře +20 % od roku 4", "Navýšení licencí/infrastruktury při překročení tieru"),
        ("A.11", "Infrastruktura / hosting privátního cloudu", "Pouze pokud ji dodává dodavatel; on-premise HW uveďte jako sizing v FS kap. 3"),
        ("A.12", "Ostatní licenční a infrastrukturní náklady", "Specifikujte v poznámce"),
    ]
    wsA, trA, offA = rocni_list("A_Licence_Infra", "A – Licence a infrastruktura", A,
                                ("Způsob zajištění", ["V ceně (OEM)", "Nakupuje Banka", "Nepotřebné", "Zahrnuto jinde"]))

    # B – implementace: MD × sazba, rok plnění
    b = wb.create_sheet("B_Implementace")
    titul(b, "B – Implementace a jednorázové služby")
    hlavicka(b, 5, ["Č.", "Položka", "Co musí obsahovat", "Počet MD", "Průměrná sazba (Kč/MD)", "Cena celkem", "Rok plnění (R1/R2)",
                    "Rok 1", "Rok 2", "Rok 3", "Rok 4", "Rok 5", "Poznámka dodavatele"])
    sirky(b, [6, 34, 50, 10, 14, 15, 10, 14, 14, 10, 10, 10, 34])
    B = [
        ("B.1", "Projektové řízení", "Po celou dobu implementace a hypercare"),
        ("B.2", "Analýza a detailní návrh řešení", "Včetně workshopů s IT, Compliance, Legal a Obchodem"),
        ("B.3", "Konfigurace produktu", "Rizikový model, scénáře TM, screening, workflow"),
        ("B.4", "Custom vývoj", "Všechny požadavky označené v Matici jako Custom vývoj"),
        ("B.5", "Integrace – CRM (REST, události)", "INT-001, KYC-011, CAS-003"),
        ("B.6", "Integrace – core banking (Kafka)", "INT-002"),
        ("B.7", "Integrace – platební systém", "INT-004"),
        ("B.8", "Integrace – DWH, DMS, ESB, registry", "INT-003, INT-005, INT-006, INT-010"),
        ("B.9", "Integrace – SIEM, monitoring, zálohování", "INT-007, INT-008; napojení na zálohovací řešení Banky"),
        ("B.10", "Datová migrace – mapování, ETL, zkušební běhy", "MIG-001, MIG-003"),
        ("B.11", "Datová migrace – rekonciliace a paralelní provoz", "MIG-002, MIG-004, min. 4 týdny paralelního běhu"),
        ("B.12", "Podpora testování (SIT, UAT) a výkonnostní testy", "NFR-003, NFR-009"),
        ("B.13", "Školení", "Train-the-trainer i koncoví uživatelé (Compliance, obchod, podpora, administrátoři)"),
        ("B.14", "Dokumentace v češtině a angličtině", "Uživatelská, administrátorská, provozní, bezpečnostní dokumentace"),
        ("B.15", "Hypercare po go-live", "Min. 3 měsíce zvýšené podpory"),
        ("B.16", "Součinnost s penetračním testem a DPIA", "SEC-006, REG-004"),
    ]
    for i, (cid, nm, desc) in enumerate(B):
        r = 6 + i
        b.cell(row=r, column=1, value=cid)
        b.cell(row=r, column=2, value=nm).alignment = WRAP
        b.cell(row=r, column=3, value=desc).alignment = WRAP
        for col in (4, 5, 7, 13):
            c = b.cell(row=r, column=col)
            c.fill = F_VEND
            c.protection = UNLOCK
        b.cell(row=r, column=4).number_format = "#,##0.0"
        b.cell(row=r, column=5).number_format = CZK
        c = b.cell(row=r, column=6, value=f"=D{r}*E{r}")
        c.fill = F_CALC
        c.number_format = CZK
        b.cell(row=r, column=8, value=f'=IF(G{r}="R1",F{r},IF(G{r}="R1+R2",F{r}/2,0))')
        b.cell(row=r, column=9, value=f'=IF(G{r}="R2",F{r},IF(G{r}="R1+R2",F{r}/2,0))')
        for col in range(10, 13):
            b.cell(row=r, column=col, value=0)
        for col in range(8, 13):
            b.cell(row=r, column=col).fill = F_CALC
            b.cell(row=r, column=col).number_format = CZK
        for col in range(1, 14):
            b.cell(row=r, column=col).border = BOX
    dv_list(b, ["R1", "R2", "R1+R2"], f"G6:G{5 + len(B)}", "R1 = rok 1, R2 = rok 2, R1+R2 = rovným dílem")
    trB = 6 + len(B)
    b.cell(row=trB, column=2, value="Mezisoučet").font = BOLD
    for col in [4, 6, 8, 9, 10, 11, 12]:
        cl = get_column_letter(col)
        c = b.cell(row=trB, column=col, value=f"=SUM({cl}6:{cl}{trB - 1})")
        c.font = BOLD
        c.fill = F_CALC
        c.number_format = CZK if col != 4 else "#,##0.0"
        c.border = BOX
    b.freeze_panes = "D6"
    b.protection.sheet = True
    b.protection.password = HESLO

    # D – rate card
    d = wb.create_sheet("D_Rate_card")
    titul(d, "D – Závazná rate card a inflační doložka")
    hlavicka(d, 5, ["Role", "Mix pro rozvoj (Banka)", "Sazba rok 1 (Kč/MD)", "Rok 2", "Rok 3", "Rok 4", "Rok 5", "Poznámka"])
    sirky(d, [34, 14, 16, 14, 14, 14, 14, 30])
    role = [("Projektový manažer", 0.10), ("Business analytik", 0.15), ("Solution architekt", 0.10),
            ("Vývojář", 0.30), ("Tester", 0.15), ("Datový specialista (migrace)", 0.05),
            ("AML doménový konzultant", 0.05), ("Specialista provozu / DevOps", 0.05), ("Bezpečnostní specialista", 0.05)]
    d["A17"] = "Navrhovaná roční indexace sazeb (dodavatel)"
    d["C17"].fill = F_VEND
    d["C17"].protection = UNLOCK
    d["C17"].number_format = PCT
    d["A18"] = "Uplatněná indexace = MIN(navrhovaná; strop 5 %)"
    d["C18"] = "=MIN(C17,0.05)"
    d["C18"].fill = F_CALC
    d["C18"].number_format = PCT
    for i, (rl, mix) in enumerate(role):
        r = 6 + i
        d.cell(row=r, column=1, value=rl)
        c = d.cell(row=r, column=2, value=mix)
        c.number_format = "0%"
        c.fill = F_BANK
        c = d.cell(row=r, column=3)
        c.fill = F_VEND
        c.protection = UNLOCK
        c.number_format = CZK
        for y in range(1, 5):
            prev = get_column_letter(2 + y)
            c = d.cell(row=r, column=3 + y, value=f"={prev}{r}*(1+$C$18)")
            c.fill = F_CALC
            c.number_format = CZK
        n = d.cell(row=r, column=8)
        n.fill = F_VEND
        n.protection = UNLOCK
        for col in range(1, 9):
            d.cell(row=r, column=col).border = BOX
    rr = 6 + len(role)
    d.cell(row=rr, column=1, value="Vážená průměrná sazba (blended)").font = BOLD
    d.cell(row=rr, column=2, value=f"=SUM(B6:B{rr - 1})").number_format = "0%"
    for y in range(5):
        cl = get_column_letter(3 + y)
        c = d.cell(row=rr, column=3 + y, value=f"=SUMPRODUCT($B$6:$B${rr - 1},{cl}6:{cl}{rr - 1})")
        c.font = BOLD
        c.fill = F_CALC
        c.number_format = CZK
        c.border = BOX
    BLEND_ROW = rr
    d.protection.sheet = True
    d.protection.password = HESLO

    # C – provoz a podpora
    C = [
        ("C.1", "Maintenance a podpora L2/L3 dle SLA (Příloha 4)", "Včetně 24/7 pro P1 a P2 incidenty"),
        ("C.2", "Upgrady a nové verze produktu", "Včetně migrace konfigurace na novou verzi"),
        ("C.3", "Provozní služby (managed service), pokud jsou nabízeny", "Jinak 0"),
        ("C.4", "Hosting privátního cloudu – provozní poplatek", "Jinak 0"),
        ("C.5", "Aktualizace dokumentace při každém releasu", "CZ/EN"),
        ("C.6", "Ostatní opakované náklady", "Specifikujte v poznámce"),
    ]
    wsC, trC, offC = rocni_list("C_Provoz_Podpora", "C – Provoz, podpora a rozvoj", C)
    # rozvojová kapacita (vzorec z rate card)
    rk = trC
    wsC.insert_rows(rk)
    wsC.cell(row=rk, column=1, value="C.7")
    wsC.cell(row=rk, column=2, value="Rozvojová kapacita (change requesty)").alignment = WRAP
    wsC.cell(row=rk, column=3, value="Vypočteno: 120 MD/rok × blended sazba z listu D (needitovat)").alignment = WRAP
    for y in range(5):
        cl = get_column_letter(3 + y)
        c = wsC.cell(row=rk, column=offC + 1 + y, value=f"={KAPACITA}*'D_Rate_card'!{cl}{BLEND_ROW}")
        c.fill = F_CALC
        c.number_format = CZK
    a_, b_ = get_column_letter(offC + 1), get_column_letter(offC + 5)
    c = wsC.cell(row=rk, column=offC + 6, value=f"=SUM({a_}{rk}:{b_}{rk})")
    c.fill = F_CALC
    c.number_format = CZK
    for col in range(1, offC + 8):
        wsC.cell(row=rk, column=col).border = BOX
    trC = rk + 1
    for col in range(offC + 1, offC + 7):
        cl = get_column_letter(col)
        wsC.cell(row=trC, column=col, value=f"=SUM({cl}6:{cl}{trC - 1})")
    # potvrzení
    q0 = trC + 2
    wsC.cell(row=q0, column=2, value="Potvrzení dodavatele (vyřazovací)").font = Font(bold=True, color=NAVY)
    potv = [
        "Legislativní údržba (AML zákon, vyhlášky ČNB, AMLR/AMLA) je zahrnuta v C.1 bez dalších poplatků",
        "Oprava kritických a vysokých zranitelností v dodaném kódu je zahrnuta v C.1 bez dalších poplatků",
        "Non-PROD prostředí (DEV, SIT, UAT, PreProd) jsou oceněna v listu A",
        "Ceny nezávisí na kurzu ani na neuvedených předpokladech",
    ]
    for i, t in enumerate(potv):
        r = q0 + 1 + i
        wsC.cell(row=r, column=2, value=t).alignment = WRAP
        wsC.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = wsC.cell(row=r, column=offC + 1)
        c.fill = F_VEND
        c.protection = UNLOCK
        dv_list(wsC, ["Ano", "Ne"], f"{get_column_letter(offC + 1)}{r}")
        wsC.row_dimensions[r].height = 30
    POTV = f"'C_Provoz_Podpora'!${get_column_letter(offC + 1)}${q0 + 1}:${get_column_letter(offC + 1)}${q0 + len(potv)}"

    E = [
        ("E.1", "Součinnost při ukončení smlouvy", "Min. 12 měsíců exitové součinnosti dle LEG-006 – ocenit v roce 5"),
        ("E.2", "Export všech dat v otevřeném strojově čitelném formátu", "Kompletní data, konfigurace, auditní stopy, dokumentace datového modelu"),
        ("E.3", "Prodloužení licencí po dobu přechodu (6 měsíců)", "Pro souběh s novým řešením"),
    ]
    wsE, trE, offE = rocni_list("E_Exit", "E – Exit a ukončení (ochrana proti vendor lock-in)", E)

    # Souhrn
    s = wb.create_sheet("Souhrn", 1)
    titul(s, "Souhrn TCO na 5 let (vypočteno – needitovat)")
    s["A4"] = "Název dodavatele:"
    s["C4"].fill = F_VEND
    s["C4"].protection = UNLOCK
    s.merge_cells("C4:E4")
    hlavicka(s, 6, ["Kategorie", ""] + ROKY + ["Celkem 5 let", "Podíl"])
    sirky(s, [44, 2, 15, 15, 15, 15, 15, 17, 9])
    zdroje = [
        ("A – Licence a infrastruktura", "A_Licence_Infra", trA, offA),
        ("B – Implementace a jednorázové služby", "B_Implementace", trB, 7),
        ("C – Provoz, podpora a rozvoj", "C_Provoz_Podpora", trC, offC),
        ("E – Exit", "E_Exit", trE, offE),
    ]
    for i, (lbl, sh, tr, off) in enumerate(zdroje):
        r = 7 + i
        s.cell(row=r, column=1, value=lbl)
        for y in range(5):
            cl = get_column_letter(off + 1 + y)
            c = s.cell(row=r, column=3 + y, value=f"='{sh}'!{cl}{tr}")
            c.number_format = CZK
        c = s.cell(row=r, column=8, value=f"=SUM(C{r}:G{r})")
        c.number_format = CZK
        c = s.cell(row=r, column=9, value=f"=IFERROR(H{r}/$H$11,0)")
        c.number_format = PCT
    s.cell(row=11, column=1, value="TCO celkem za 5 let (HODNOCENÁ CENA)").font = BOLD
    for col in range(3, 9):
        cl = get_column_letter(col)
        c = s.cell(row=11, column=col, value=f"=SUM({cl}7:{cl}10)")
        c.font = BOLD
        c.number_format = CZK
    s.cell(row=12, column=1, value="NPV TCO (informativně, nehodnotí se)")
    c = s.cell(row=12, column=8, value=f"=NPV({DISK},C11:G11)")
    c.number_format = CZK
    s.cell(row=13, column=1, value="TCO v EUR (informativně)")
    c = s.cell(row=13, column=8, value=f"=H11/{KURZ}")
    c.number_format = '#,##0 "EUR"'
    for row in s.iter_rows(min_row=7, max_row=13, max_col=9):
        for c in row:
            c.border = BOX
            if c.column > 2:
                c.fill = F_CALC
    s.cell(row=15, column=1, value="Kontroly úplnosti").font = Font(bold=True, color=NAVY, size=12)
    chk = [
        ("Nevyplněné cenové buňky v listu A", "=COUNTBLANK('A_Licence_Infra'!E6:I17)"),
        ("Nevyplněný způsob zajištění v listu A", "=COUNTBLANK('A_Licence_Infra'!D6:D17)"),
        ("Nevyplněné MD / sazby / rok plnění v listu B", f"=COUNTBLANK('B_Implementace'!D6:E{trB - 1})+COUNTBLANK('B_Implementace'!G6:G{trB - 1})"),
        ("Nevyplněné cenové buňky v listu C", f"=COUNTBLANK('C_Provoz_Podpora'!D6:H{6 + len(C) - 1})"),
        ("Nevyplněné sazby v rate card / indexace", f"=COUNTBLANK('D_Rate_card'!C6:C{BLEND_ROW - 1})+COUNTBLANK('D_Rate_card'!C17)"),
        ("Nevyplněné cenové buňky v listu E", f"=COUNTBLANK('E_Exit'!D6:H{6 + len(E) - 1})"),
        ("Potvrzení na listu C s hodnotou jinou než Ano", f'=COUNTA({POTV})-COUNTIF({POTV},"Ano")+COUNTBLANK({POTV})'),
    ]
    for i, (lbl, f) in enumerate(chk):
        r = 16 + i
        s.cell(row=r, column=1, value=lbl)
        c = s.cell(row=r, column=3, value=f)
        c.border = BOX
        c.fill = F_CALC
    s.cell(row=16 + len(chk), column=1, value="Stav nabídky").font = BOLD
    c = s.cell(row=16 + len(chk), column=3, value=f'=IF(SUM(C16:C{15 + len(chk)})=0,"KOMPLETNÍ","NEÚPLNÁ")')
    c.font = BOLD
    s.merge_cells(start_row=16 + len(chk), start_column=3, end_row=16 + len(chk), end_column=4)
    s.conditional_formatting.add(f"C16:C{15 + len(chk)}", CellIsRule(operator="greaterThan", formula=["0"], fill=F_RED))
    s.conditional_formatting.add(f"C16:C{15 + len(chk)}", CellIsRule(operator="equal", formula=["0"], fill=F_GREEN))
    s.conditional_formatting.add(f"C{16 + len(chk)}", CellIsRule(operator="equal", formula=['"KOMPLETNÍ"'], fill=F_GREEN))
    s.conditional_formatting.add(f"C{16 + len(chk)}", CellIsRule(operator="equal", formula=['"NEÚPLNÁ"'], fill=F_RED))
    s.protection.sheet = True
    s.protection.password = HESLO
    # pořadí listů
    order = ["Návod", "Souhrn", "Předpoklady", "A_Licence_Infra", "B_Implementace", "C_Provoz_Podpora", "D_Rate_card", "E_Exit"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    wb.save(path)


# ---------------------------------------------------------------- Hodnoticí model
DODAVATELE = ["Dodavatel A", "Dodavatel B", "Dodavatel C"]
KOMISE = [
    ("Ing. Jana Novotná", "Business analytička / Product Owner AML", "předsedkyně, hlasuje"),
    ("Ing. Martin Dvořák", "Enterprise architekt (IT)", "hlasuje"),
    ("Mgr. Tomáš Svoboda", "Bezpečnostní architekt (útvar CISO)", "hlasuje"),
    ("JUDr. Lucie Malá", "Ředitelka Compliance, pověřená osoba (MLRO)", "hlasuje"),
    ("Ing. Pavel Černý", "Procurement", "bez hlasovacího práva, dohled nad procesem"),
]


def hodnotici_model(path):
    rnd = random.Random(2506)
    wb = Workbook()
    navod(wb, "Hodnoticí model a hodnoticí arch – výběrové řízení SENTINEL", [
        "# Postup hodnocení",
        "1) K.O. brána (list 1_KO_brana): každé K.O. kritérium Pass/Fail. Jediný FAIL = nabídka nepostupuje do bodového hodnocení.",
        "2) Byznys funkcionality 30 % (list 2_Byznys): odpovědi z Matice shody, body 5/3/1/0, váha MoSCoW, křížová kontrola s FS a ověření v demu.",
        "3) Architektura a bezpečnost 20 % (list 3_Arch_Sec): každý hlasující člen komise boduje kapitoly FS na škále 0–10.",
        "4) Delivery a zkušenosti 10 % (list 4_Delivery): reference z finančního sektoru, seniorita týmu, návrh SLA, projektový plán – škála 0–10.",
        "5) Cena 40 % (list 5_Cena): (nejnižší TCO / hodnocené TCO) × 40, jen pro nabídky, které prošly K.O. bránou.",
        "6) Výsledek (list Vysledek) a hodnoticí arch k podpisu (list Hodnotici_arch).",
        "# Pravidla auditní stopy",
        "• Váhy a stupnice schválil SteerCo 23. 9. 2026, tj. PŘED vyhlášením výběrového řízení. Po otevření nabídek se nemění.",
        "• Každá známka horší než průměr daného kritéria musí mít jednovětné zdůvodnění – buňka jinak zčervená.",
        "• Členové komise podepisují prohlášení o neexistenci střetu zájmů před otevřením nabídek.",
        "• Data v tomto souboru jsou FIKTIVNÍ ukázkou výpočtu (3 smyšlení dodavatelé).",
        "# Barvy",
    ])
    legenda(wb.active, 19)

    par = wb.create_sheet("Parametry")
    titul(par, "Parametry hodnocení (schváleno SteerCo 23. 9. 2026)")
    hlavicka(par, 5, ["Kategorie", "Váha", "Zdroj dat", "Metoda"])
    sirky(par, [30, 9, 34, 70])
    kat = [
        ("Cena (TCO na 5 let)", 0.40, "Příloha 3 – cenová šablona", "(nejnižší TCO / hodnocené TCO) × 40"),
        ("Byznys funkcionality", 0.30, "Příloha 1 – Matice shody", "Σ(uznané body × váha MoSCoW) / Σ(5 × váha MoSCoW) × 30"),
        ("Architektura & bezpečnost", 0.20, "Příloha 2 – Feasibility Study", "průměr známek komise 0–10 / 10 × 20"),
        ("Delivery a zkušenosti", 0.10, "Reference, CV, návrh SLA, plán", "průměr známek komise 0–10 / 10 × 10"),
    ]
    for i, row in enumerate(kat):
        for j, v in enumerate(row):
            c = par.cell(row=6 + i, column=j + 1, value=v)
            c.border = BOX
            c.alignment = WRAP
            c.fill = F_BANK
        par.cell(row=6 + i, column=2).number_format = "0%"
    par["A10"] = "Součet vah"
    par["B10"] = "=SUM(B6:B9)"
    par["B10"].number_format = "0%"
    par["A10"].font = BOLD
    W = {"cena": "Parametry!$B$6", "byz": "Parametry!$B$7", "arch": "Parametry!$B$8", "del": "Parametry!$B$9"}
    par["A12"] = "Stupnice FS a Delivery (0–10)"
    par["A12"].font = BOLD
    for i, t in enumerate(["8–10: detailní návrh přesně na rozhraní a systémy Banky",
                           "4–7: standardizovaný popis, chybí detail na specifické systémy Banky (core, CRM)",
                           "0–3: obecné marketingové fráze bez technické hloubky"]):
        par.cell(row=13 + i, column=1, value=t)

    # 1 KO brána
    ko = wb.create_sheet("1_KO_brana")
    titul(ko, "1 – K.O. brána (Pass/Fail)")
    hlavicka(ko, 5, ["ID", "K.O. kritérium", "Doklad"] + DODAVATELE + ["Poznámka komise"],
             [F_HEAD, F_HEAD, F_HEAD, F_COMM, F_COMM, F_COMM, F_HEAD])
    sirky(ko, [10, 60, 36, 13, 13, 13, 60])
    ko_rows = [
        ("FORM-01", "Nabídka podána včas a v požadované struktuře (Matice, FS, TCO, přílohy)", "Protokol o otevření nabídek"),
        ("FORM-02", "Podepsaná NDA a čestné prohlášení (bezúhonnost, sankce, střet zájmů)", "Čestné prohlášení dle Přílohy 7"),
        ("FORM-03", "Kompletní cenová šablona – stav KOMPLETNÍ a všechna potvrzení „Ano“", "Příloha 3, list Souhrn"),
    ] + [(r[0], r[2], f"Matice shody, FS kap. {r[8]}") for r in P.POZADAVKY if r[6] == "Ano"]
    fails = {("Dodavatel C", "NFR-006"): "Zálohy v regionu mimo EU/EHP (USA); 2. linie podpory z Indie.",
             ("Dodavatel C", "LEG-001"): "Neakceptuje omezení vzdáleného přístupu z třetích zemí.",
             }
    for i, (rid, txt, dok) in enumerate(ko_rows):
        r = 6 + i
        ko.cell(row=r, column=1, value=rid)
        ko.cell(row=r, column=2, value=txt).alignment = WRAP
        ko.cell(row=r, column=3, value=dok).alignment = WRAP
        poz = []
        for j, dname in enumerate(DODAVATELE):
            v = "FAIL" if (dname, rid) in fails else "PASS"
            if (dname, rid) in fails:
                poz.append(f"{dname}: {fails[(dname, rid)]}")
            c = ko.cell(row=r, column=4 + j, value=v)
            c.fill = F_COMM
            c.alignment = CENTER
        ko.cell(row=r, column=7, value=" ".join(poz)).alignment = WRAP
        for col in range(1, 8):
            ko.cell(row=r, column=col).border = BOX
    kr = 6 + len(ko_rows)
    dv_list(ko, ["PASS", "FAIL"], f"D6:F{kr - 1}")
    ko.cell(row=kr, column=2, value="Výsledek K.O. brány").font = BOLD
    for j in range(3):
        cl = get_column_letter(4 + j)
        c = ko.cell(row=kr, column=4 + j, value=f'=IF(COUNTIF({cl}6:{cl}{kr - 1},"FAIL")>0,"VYŘAZEN",IF(COUNTIF({cl}6:{cl}{kr - 1},"PASS")={kr - 6},"POSTUPUJE","NEÚPLNÉ"))')
        c.font = BOLD
        c.alignment = CENTER
        c.border = BOX
    ko.conditional_formatting.add(f"D6:F{kr}", CellIsRule(operator="equal", formula=['"FAIL"'], fill=F_RED))
    ko.conditional_formatting.add(f"D6:F{kr}", CellIsRule(operator="equal", formula=['"VYŘAZEN"'], fill=F_RED))
    ko.conditional_formatting.add(f"D6:F{kr}", CellIsRule(operator="equal", formula=['"POSTUPUJE"'], fill=F_GREEN))
    KO_RES = {d_: f"'1_KO_brana'!{get_column_letter(4 + j)}{kr}" for j, d_ in enumerate(DODAVATELE)}

    # 2 Byznys
    by = wb.create_sheet("2_Byznys")
    titul(by, "2 – Byznys funkcionality (30 %) – odpovědi z Matice shody")
    base = ["ID", "Požadavek", "MoSCoW", "Váha"]
    heads = list(base)
    for dname in DODAVATELE:
        heads += [f"{dname}\nOdpověď", "Odkaz FS", "Ověřeno (FS/demo)", "Uznané body"]
    heads += ["Zdůvodnění komise (povinné při snížení bodů)"]
    fills = [F_HEAD] * 4 + [F_VEND, F_VEND, F_COMM, F_CALC] * 3 + [F_COMM]
    hlavicka(by, 5, heads, fills)
    by.row_dimensions[5].height = 42
    sirky(by, [10, 52, 8, 6] + [13, 8, 10, 9] * 3 + [50])
    prof = {
        "Dodavatel A": [("OOTB", .58), ("Konfigurace", .32), ("Custom vývoj", .08), ("Nesplněno", .02)],
        "Dodavatel B": [("OOTB", .32), ("Konfigurace", .36), ("Custom vývoj", .24), ("3rd party", .08)],
        "Dodavatel C": [("OOTB", .62), ("Konfigurace", .28), ("Custom vývoj", .06), ("Nesplněno", .04)],
    }

    def pick(dname):
        x = rnd.random()
        acc = 0
        for o, pr in prof[dname]:
            acc += pr
            if x <= acc:
                return o
        return prof[dname][-1][0]

    byz = [r for r in P.POZADAVKY if r[1] in P.BYZNYS and r[6] != "Ano"]
    special = {("Dodavatel B", "TM-002"): ("Custom vývoj", "", "Ano", "Chybí odkaz do FS – dle pravidla křížové kontroly nesplněno."),
               ("Dodavatel B", "SCR-005"): ("OOTB", "6.2", "Ne", "V demu transliterace cyrilice nefungovala, shoda nalezena jen pro latinku."),
               ("Dodavatel A", "TM-007"): ("Custom vývoj", "6.6", "Ano", ""),
               ("Dodavatel A", "KYC-004"): ("OOTB", "6.1", "Ano", ""),
               }
    first = 6
    for i, req in enumerate(byz):
        r = first + i
        by.cell(row=r, column=1, value=req[0])
        by.cell(row=r, column=2, value=req[2]).alignment = WRAP
        by.cell(row=r, column=3, value=req[5]).alignment = CENTER
        by.cell(row=r, column=4, value=f'=IF(C{r}="Must",3,IF(C{r}="Should",2,IF(C{r}="Could",1,0)))').alignment = CENTER
        notes = []
        for j, dname in enumerate(DODAVATELE):
            c0 = 5 + j * 4
            if (dname, req[0]) in special:
                odp, fs, ov, nt = special[(dname, req[0])]
                if nt:
                    notes.append(f"{dname}: {nt}")
            else:
                odp = pick(dname)
                fs = req[8] if odp != "Nesplněno" and odp != "3rd party" else ""
                if odp == "Custom vývoj":
                    fs = "6.6"
                ov = "Ano" if odp not in ("Nesplněno", "3rd party") else ""
            O, F_, V = (get_column_letter(c0 + k) for k in range(3))
            by.cell(row=r, column=c0, value=odp)
            by.cell(row=r, column=c0 + 1, value=fs)
            by.cell(row=r, column=c0 + 2, value=ov)
            by.cell(row=r, column=c0 + 3, value=(
                f'=IF({O}{r}="","",IF(OR(AND(OR({O}{r}="OOTB",{O}{r}="Konfigurace",{O}{r}="Custom vývoj"),{F_}{r}=""),{V}{r}="Ne"),0,'
                f'IF({O}{r}="OOTB",5,IF({O}{r}="Konfigurace",3,IF({O}{r}="Custom vývoj",1,0)))))'))
            for k, fl in enumerate([F_VEND, F_VEND, F_COMM, F_CALC]):
                cc = by.cell(row=r, column=c0 + k)
                cc.fill = fl
                cc.alignment = CENTER
        by.cell(row=r, column=17, value=" ".join(notes)).alignment = WRAP
        for col in range(1, 18):
            by.cell(row=r, column=col).border = BOX
    last = first + len(byz) - 1
    for j in range(3):
        c0 = 5 + j * 4
        dv_list(by, ODPOVEDI, f"{get_column_letter(c0)}{first}:{get_column_letter(c0)}{last}")
        dv_list(by, ["Ano", "Ne"], f"{get_column_letter(c0 + 2)}{first}:{get_column_letter(c0 + 2)}{last}")
    sr = last + 2
    by.cell(row=sr, column=2, value="Pokrytí byznys funkcionalit (vážené)").font = BOLD
    by.cell(row=sr + 1, column=2, value="Body (max. 30)").font = BOLD
    by.cell(row=sr + 2, column=2, value="Počet odpovědí Custom vývoj").font = BOLD
    by.cell(row=sr + 3, column=2, value="Počet nesplněných / 3rd party").font = BOLD
    BYZ = {}
    for j, dname in enumerate(DODAVATELE):
        c0 = 5 + j * 4
        U = get_column_letter(c0 + 3)
        O = get_column_letter(c0)
        c = by.cell(row=sr, column=c0 + 3, value=f"=SUMPRODUCT(N(+{U}{first}:{U}{last}),$D${first}:$D${last})/(5*SUM($D${first}:$D${last}))")
        c.number_format = PCT
        c = by.cell(row=sr + 1, column=c0 + 3, value=f"={U}{sr}*{W['byz']}*100")
        c.number_format = "0.00"
        by.cell(row=sr + 2, column=c0 + 3, value=f'=COUNTIF({O}{first}:{O}{last},"Custom vývoj")')
        by.cell(row=sr + 3, column=c0 + 3, value=f'=COUNTIF({O}{first}:{O}{last},"Nesplněno")+COUNTIF({O}{first}:{O}{last},"3rd party")')
        for k in range(4):
            by.cell(row=sr + k, column=c0 + 3).fill = F_CALC
            by.cell(row=sr + k, column=c0 + 3).font = BOLD
        BYZ[dname] = f"'2_Byznys'!{U}{sr + 1}"
    by.freeze_panes = "C6"
    by.conditional_formatting.add(f"Q{first}:Q{last}", FormulaRule(
        formula=[f'AND(Q{first}="",OR(AND(ISNUMBER(H{first}),H{first}=0,E{first}<>"Nesplněno",E{first}<>"3rd party"),'
                 f'AND(ISNUMBER(L{first}),L{first}=0,I{first}<>"Nesplněno",I{first}<>"3rd party"),'
                 f'AND(ISNUMBER(P{first}),P{first}=0,M{first}<>"Nesplněno",M{first}<>"3rd party")))'], fill=F_RED))

    # 3 Arch & Sec
    ar = wb.create_sheet("3_Arch_Sec")
    titul(ar, "3 – Architektura a bezpečnost (20 %) – hodnocení Feasibility Study, škála 0–10")
    hlasujici = [k for k in KOMISE if "hlasuje" in k[2]]
    kriteria = [
        ("FS 3", "Návrh cílové architektury (hosting, DB, HA/DR, sizing)"),
        ("FS 4", "Bezpečnostní koncept (role, šifrování, logování, zranitelnosti, pentesty)"),
        ("FS 5", "Řešení kritických integrací (CRM, core přes Kafka, platby, DWH)"),
        ("FS 6.6", "Řešení custom požadavků"),
        ("FS 7", "Plán implementace a migrace bez výpadku"),
        ("FS 9–10", "Regulatorní soulad, auditovatelnost a exit strategie"),
    ]
    skore = {
        "Dodavatel A": [[9, 8, 8, 8], [8, 9, 8, 8], [8, 8, 9, 8], [7, 7, 8, 7], [8, 8, 8, 9], [8, 8, 9, 9]],
        "Dodavatel B": [[7, 7, 7, 7], [6, 5, 6, 7], [5, 6, 6, 6], [6, 6, 6, 5], [7, 7, 6, 7], [6, 6, 7, 6]],
        "Dodavatel C": [[8, 7, 8, 8], [7, 6, 7, 7], [8, 7, 8, 8], [8, 8, 8, 8], [7, 7, 7, 7], [5, 4, 5, 4]],
    }
    zduv = {
        ("Dodavatel B", "FS 5"): "Dodavatel navrhuje point-to-point integraci na core banking namísto požadované bankovní Kafka sběrnice.",
        ("Dodavatel B", "FS 4"): "Správa šifrovacích klíčů u dodavatele, ne v HSM Banky; chybí harmonogram penetračních testů.",
        ("Dodavatel B", "FS 6.6"): "Custom vývoj popsán obecně bez odhadu pracnosti a dopadu na upgrady.",
        ("Dodavatel B", "FS 9–10"): "Exit plán neřeší export auditních stop a historie screeningu.",
        ("Dodavatel C", "FS 9–10"): "Zálohy mimo EU/EHP a nejasné subdodavatelské řetězení podpory; exit bez exportu konfigurace.",
        ("Dodavatel C", "FS 4"): "Privilegovaný přístup podpory z třetí země není řešen přes PAM Banky.",
        ("Dodavatel B", "FS 3"): "Sizing HW uveden jen pro PROD, chybí PreProd a DR.",
        ("Dodavatel B", "FS 7"): "Paralelní provoz jen 2 týdny a bez popsané rekonciliace historie případů.",
        ("Dodavatel C", "FS 3"): "Architektura počítá se sdílenou multi-tenant platformou, privátní instance jen jako příplatková varianta.",
        ("Dodavatel C", "FS 5"): "Integrace na CRM popsána obecně přes standardní konektor, chybí řešení událostí pro změny stavu klienta.",
        ("Dodavatel C", "FS 7"): "Migrace historie screeningu jen za 5 let, zbytek jako archiv v PDF – nesplňuje § 16 AML zákona.",
    }
    hlavicka(ar, 5, ["Dodavatel", "Kritérium (kapitola FS)", "Popis"] + [k[0] for k in hlasujici] + ["Průměr", "Průměr kritéria (všichni dodavatelé)", "Zdůvodnění (povinné u známky pod průměrem kritéria)", "Kontrola"],
             [F_HEAD] * 3 + [F_COMM] * len(hlasujici) + [F_CALC, F_CALC, F_COMM, F_CALC])
    ar.row_dimensions[5].height = 42
    sirky(ar, [13, 12, 44] + [12] * len(hlasujici) + [9, 13, 60, 16])
    r = 6
    ARCH = {}
    nh = len(hlasujici)
    sc, ec = get_column_letter(4), get_column_letter(3 + nh)
    avg_c = 4 + nh
    krit_rows = {k[0]: [] for k in kriteria}
    vend_rows = {}
    for dname in DODAVATELE:
        rows_v = []
        for ki, (kid, kdesc) in enumerate(kriteria):
            ar.cell(row=r, column=1, value=dname)
            ar.cell(row=r, column=2, value=kid)
            ar.cell(row=r, column=3, value=kdesc).alignment = WRAP
            for m in range(nh):
                c = ar.cell(row=r, column=4 + m, value=skore[dname][ki][m])
                c.fill = F_COMM
                c.alignment = CENTER
            c = ar.cell(row=r, column=avg_c, value=f"=AVERAGE({sc}{r}:{ec}{r})")
            c.number_format = "0.00"
            c.fill = F_CALC
            ar.cell(row=r, column=avg_c + 2, value=zduv.get((dname, kid), "")).alignment = WRAP
            krit_rows[kid].append(r)
            rows_v.append(r)
            for col in range(1, avg_c + 4):
                ar.cell(row=r, column=col).border = BOX
            r += 1
        vend_rows[dname] = rows_v
    for kid, rws in krit_rows.items():
        rng_ = ",".join(f"{sc}{x}:{ec}{x}" for x in rws)
        for x in rws:
            c = ar.cell(row=x, column=avg_c + 1, value=f"=AVERAGE({rng_})")
            c.number_format = "0.00"
            c.fill = F_CALC
            Z = get_column_letter(avg_c + 2)
            Pm = get_column_letter(avg_c + 1)
            c = ar.cell(row=x, column=avg_c + 3, value=f'=IF(AND(MIN({sc}{x}:{ec}{x})<{Pm}{x},{Z}{x}=""),"DOPLNIT ZDŮVODNĚNÍ","OK")')
            c.fill = F_CALC
    dv_list(ar, [str(i) for i in range(11)], f"{sc}6:{ec}{r - 1}")
    ar.conditional_formatting.add(f"{get_column_letter(avg_c + 3)}6:{get_column_letter(avg_c + 3)}{r - 1}",
                                  CellIsRule(operator="equal", formula=['"DOPLNIT ZDŮVODNĚNÍ"'], fill=F_RED))
    r += 1
    ar.cell(row=r, column=1, value="Výsledek").font = Font(bold=True, color=NAVY, size=12)
    hlavicka(ar, r + 1, ["Dodavatel", "Průměrná známka", "Body (max. 20)"])
    A_ = get_column_letter(avg_c)
    for j, dname in enumerate(DODAVATELE):
        rr = r + 2 + j
        rws = vend_rows[dname]
        ar.cell(row=rr, column=1, value=dname)
        c = ar.cell(row=rr, column=2, value=f"=AVERAGE({A_}{rws[0]}:{A_}{rws[-1]})")
        c.number_format = "0.00"
        c = ar.cell(row=rr, column=3, value=f"=B{rr}/10*{W['arch']}*100")
        c.number_format = "0.00"
        c.font = BOLD
        ARCH[dname] = f"'3_Arch_Sec'!C{rr}"
    ar.freeze_panes = "D6"

    # 4 Delivery
    de = wb.create_sheet("4_Delivery")
    titul(de, "4 – Delivery a zkušenosti (10 %), škála 0–10")
    dkrit = [
        ("D1", "Reference z finančního sektoru (min. 2 banky v EU s AML řešením v produkci)"),
        ("D2", "Seniorita klíčových osob (CV: PM, architekt, AML konzultant, vedoucí migrace)"),
        ("D3", "Návrh SLA a podpory vs. Příloha 4 (odchylky, sankce, RCA)"),
        ("D4", "Realističnost projektového plánu a kapacit"),
    ]
    dsk = {"Dodavatel A": [[9, 8, 8, 9], [8, 8, 7, 8], [7, 7, 8, 7], [8, 8, 8, 8]],
           "Dodavatel B": [[8, 8, 8, 8], [7, 7, 7, 8], [8, 8, 7, 8], [6, 7, 6, 6]],
           "Dodavatel C": [[9, 9, 8, 9], [7, 7, 7, 7], [5, 5, 4, 5], [7, 7, 7, 7]]}
    dzd = {("Dodavatel B", "D4"): "Plán počítá s paralelním provozem jen 2 týdny místo požadovaných 4.",
           ("Dodavatel C", "D3"): "SLA měří dostupnost infrastruktury, ne byznys transakcí; cap sankcí 5 % ročního poplatku.",
           ("Dodavatel A", "D3"): "Navrhuje cap sankcí 15 % ročního poplatku, Banka požaduje min. 30 %.",
           ("Dodavatel A", "D1"): "Dvě reference z bank v EU, ale žádná z ČR – chybí přímá zkušenost s požadavky ČNB a FAÚ.",
           ("Dodavatel A", "D2"): "Vedoucí migrace má méně než 3 roky praxe s migrací AML dat.",
           ("Dodavatel B", "D1"): "Reference pouze z retailových bank, žádná z privátního bankovnictví.",
           ("Dodavatel B", "D2"): "Navržený solution architekt nemá zkušenost s integrací přes Kafka.",
           ("Dodavatel C", "D1"): "Jedna ze dvou referencí je platební instituce, nikoli banka.",
           ("Dodavatel C", "D2"): "CV dvou klíčových osob jsou od subdodavatele a neodpovídají týmu v projektovém plánu.",
           ("Dodavatel C", "D4"): "Plán nemá rezervu na UAT obchodu ani na opakované zkušební migrace."}
    hlavicka(de, 5, ["Dodavatel", "Kritérium", "Popis"] + [k[0] for k in hlasujici] + ["Průměr", "Průměr kritéria", "Zdůvodnění", "Kontrola"],
             [F_HEAD] * 3 + [F_COMM] * nh + [F_CALC, F_CALC, F_COMM, F_CALC])
    de.row_dimensions[5].height = 42
    sirky(de, [13, 8, 52] + [12] * nh + [9, 10, 60, 16])
    r = 6
    krit_rows = {k[0]: [] for k in dkrit}
    vend_rows = {}
    for dname in DODAVATELE:
        rows_v = []
        for ki, (kid, kdesc) in enumerate(dkrit):
            de.cell(row=r, column=1, value=dname)
            de.cell(row=r, column=2, value=kid)
            de.cell(row=r, column=3, value=kdesc).alignment = WRAP
            for m in range(nh):
                c = de.cell(row=r, column=4 + m, value=dsk[dname][ki][m])
                c.fill = F_COMM
                c.alignment = CENTER
            c = de.cell(row=r, column=avg_c, value=f"=AVERAGE({sc}{r}:{ec}{r})")
            c.number_format = "0.00"
            c.fill = F_CALC
            de.cell(row=r, column=avg_c + 2, value=dzd.get((dname, kid), "")).alignment = WRAP
            krit_rows[kid].append(r)
            rows_v.append(r)
            for col in range(1, avg_c + 4):
                de.cell(row=r, column=col).border = BOX
            r += 1
        vend_rows[dname] = rows_v
    for kid, rws in krit_rows.items():
        rng_ = ",".join(f"{sc}{x}:{ec}{x}" for x in rws)
        for x in rws:
            c = de.cell(row=x, column=avg_c + 1, value=f"=AVERAGE({rng_})")
            c.number_format = "0.00"
            c.fill = F_CALC
            Z = get_column_letter(avg_c + 2)
            Pm = get_column_letter(avg_c + 1)
            c = de.cell(row=x, column=avg_c + 3, value=f'=IF(AND(MIN({sc}{x}:{ec}{x})<{Pm}{x},{Z}{x}=""),"DOPLNIT ZDŮVODNĚNÍ","OK")')
            c.fill = F_CALC
    dv_list(de, [str(i) for i in range(11)], f"{sc}6:{ec}{r - 1}")
    de.conditional_formatting.add(f"{get_column_letter(avg_c + 3)}6:{get_column_letter(avg_c + 3)}{r - 1}",
                                  CellIsRule(operator="equal", formula=['"DOPLNIT ZDŮVODNĚNÍ"'], fill=F_RED))
    r += 1
    hlavicka(de, r + 1, ["Dodavatel", "Průměrná známka", "Body (max. 10)"])
    DEL = {}
    for j, dname in enumerate(DODAVATELE):
        rr = r + 2 + j
        rws = vend_rows[dname]
        de.cell(row=rr, column=1, value=dname)
        c = de.cell(row=rr, column=2, value=f"=AVERAGE({A_}{rws[0]}:{A_}{rws[-1]})")
        c.number_format = "0.00"
        c = de.cell(row=rr, column=3, value=f"=B{rr}/10*{W['del']}*100")
        c.number_format = "0.00"
        c.font = BOLD
        DEL[dname] = f"'4_Delivery'!C{rr}"
    de.freeze_panes = "D6"

    # 5 Cena
    ce = wb.create_sheet("5_Cena")
    titul(ce, "5 – Cena (40 %): TCO na 5 let z Přílohy 3 (po BAFO)")
    hlavicka(ce, 5, ["Dodavatel", "TCO 5 let (Kč bez DPH)", "Stav K.O. brány", "TCO do výpočtu", "Body (max. 40)", "Poznámka"],
             [F_HEAD, F_VEND, F_CALC, F_CALC, F_CALC, F_HEAD])
    sirky(ce, [14, 22, 16, 22, 14, 60])
    ceny = {"Dodavatel A": 52_400_000, "Dodavatel B": 41_900_000, "Dodavatel C": 46_300_000}
    cpo = {"Dodavatel A": "Po BAFO sleva 6 % na licence; datové feedy PEP oceněny na 5 let.",
           "Dodavatel B": "Nízká cena licence; 24 požadavků jako Custom vývoj zvyšuje riziko víceprací.",
           "Dodavatel C": "Vyřazen v K.O. bráně – cena se nehodnotí."}
    CEN = {}
    for j, dname in enumerate(DODAVATELE):
        r = 6 + j
        ce.cell(row=r, column=1, value=dname)
        c = ce.cell(row=r, column=2, value=ceny[dname])
        c.number_format = CZK
        c.fill = F_VEND
        ce.cell(row=r, column=3, value=f"={KO_RES[dname]}").fill = F_CALC
        c = ce.cell(row=r, column=4, value=f'=IF(C{r}="POSTUPUJE",B{r},"")')
        c.number_format = CZK
        c.fill = F_CALC
        c = ce.cell(row=r, column=5, value=f'=IF(D{r}="","",MIN($D$6:$D$8)/D{r}*{W["cena"]}*100)')
        c.number_format = "0.00"
        c.fill = F_CALC
        c.font = BOLD
        ce.cell(row=r, column=6, value=cpo[dname]).alignment = WRAP
        for col in range(1, 7):
            ce.cell(row=r, column=col).border = BOX
        CEN[dname] = f"'5_Cena'!E{r}"

    # Výsledek
    vy = wb.create_sheet("Vysledek")
    titul(vy, "Výsledek hodnocení")
    hlavicka(vy, 5, ["Dodavatel", "K.O. brána", "Cena (40)", "Byznys (30)", "Arch & Sec (20)", "Delivery (10)", "Celkem (100)", "Pořadí"])
    sirky(vy, [16, 14, 11, 11, 13, 12, 13, 9])
    for j, dname in enumerate(DODAVATELE):
        r = 6 + j
        vy.cell(row=r, column=1, value=dname).font = BOLD
        vy.cell(row=r, column=2, value=f"={KO_RES[dname]}")
        vy.cell(row=r, column=3, value=f'=IF(B{r}="POSTUPUJE",{CEN[dname]},"")')
        vy.cell(row=r, column=4, value=f'=IF(B{r}="POSTUPUJE",{BYZ[dname]},"")')
        vy.cell(row=r, column=5, value=f'=IF(B{r}="POSTUPUJE",{ARCH[dname]},"")')
        vy.cell(row=r, column=6, value=f'=IF(B{r}="POSTUPUJE",{DEL[dname]},"")')
        vy.cell(row=r, column=7, value=f'=IF(B{r}="POSTUPUJE",SUM(C{r}:F{r}),"vyřazen")')
        vy.cell(row=r, column=8, value=f'=IF(ISNUMBER(G{r}),RANK(G{r},$G$6:$G$8),"–")')
        for col in range(1, 9):
            c = vy.cell(row=r, column=col)
            c.border = BOX
            c.alignment = CENTER
            if 3 <= col <= 7:
                c.number_format = "0.00"
        vy.cell(row=r, column=7).font = BOLD
    vy.conditional_formatting.add("H6:H8", CellIsRule(operator="equal", formula=["1"], fill=F_GREEN))
    vy.conditional_formatting.add("B6:B8", CellIsRule(operator="equal", formula=['"VYŘAZEN"'], fill=F_RED))
    vy["A10"] = "Citlivost: rozdíl 1. a 2. místa v bodech"
    vy["G10"] = "=LARGE($G$6:$G$8,1)-LARGE($G$6:$G$8,2)"
    vy["G10"].number_format = "0.00"
    vy["A11"] = "Pravidlo shody bodů: při rozdílu < 1,0 bodu rozhoduje vyšší skóre Arch & Sec; poté nižší TCO."
    vy["A11"].font = SUB

    # Hodnoticí arch
    ha = wb.create_sheet("Hodnotici_arch")
    titul(ha, "Hodnoticí arch – výběrové řízení SENTINEL (k podpisu)")
    sirky(ha, [28, 40, 34, 22, 22])
    ha["A5"] = "Předmět:"
    ha["B5"] = "Dodávka, implementace a podpora KYC/AML řešení (screening, KYC, monitoring transakcí, správa případů)"
    ha["A6"] = "Datum hodnocení:"
    ha["B6"] = "2027-02-17"
    ha["A7"] = "Počet nabídek:"
    ha["B7"] = 3
    ha["A9"] = "Výsledek"
    ha["A9"].font = Font(bold=True, color=NAVY, size=12)
    hlavicka(ha, 10, ["Dodavatel", "K.O. brána", "Celkem bodů", "Pořadí"])
    for j, dname in enumerate(DODAVATELE):
        r = 11 + j
        ha.cell(row=r, column=1, value=dname)
        ha.cell(row=r, column=2, value=f"=Vysledek!B{6 + j}")
        c = ha.cell(row=r, column=3, value=f"=Vysledek!G{6 + j}")
        c.number_format = "0.00"
        ha.cell(row=r, column=4, value=f"=Vysledek!H{6 + j}")
        for col in range(1, 5):
            ha.cell(row=r, column=col).border = BOX
    ha["A15"] = "Doporučení komise:"
    ha["A15"].font = BOLD
    ha["B15"] = '=IFERROR("Doporučujeme zahájit jednání o smlouvě s: "&INDEX(Vysledek!A6:A8,MATCH(1,Vysledek!H6:H8,0)),"")'
    ha["A17"] = "Zdůvodnění horších známek (výtah z listů 2–4)"
    ha["A17"].font = BOLD
    r = 18
    for (dname, kid), t in sorted({**zduv, **dzd}.items()):
        ha.cell(row=r, column=1, value=f"{dname} – {kid}")
        ha.cell(row=r, column=2, value=t).alignment = WRAP
        ha.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ha.row_dimensions[r].height = 30
        r += 1
    r += 1
    ha.cell(row=r, column=1, value="Prohlášení: Členové komise prohlašují, že nejsou ve střetu zájmů ve vztahu k žádnému z uchazečů, "
            "hodnotili nezávisle podle předem schválené metodiky a výsledky odpovídají podkladům v tomto souboru.").alignment = WRAP
    ha.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ha.row_dimensions[r].height = 45
    r += 2
    hlavicka(ha, r, ["Jméno", "Role", "Postavení v komisi", "Datum", "Podpis"])
    for k in KOMISE:
        r += 1
        for j, v in enumerate(list(k) + ["", ""]):
            c = ha.cell(row=r, column=j + 1, value=v)
            c.border = BOX
            c.alignment = WRAP
        ha.row_dimensions[r].height = 30
    r += 2
    ha.cell(row=r, column=1, value="Schválil (sponzor projektu):")
    ha.cell(row=r, column=2, value="Ing. Karel Procházka, člen představenstva")
    ha.cell(row=r, column=4, value="Podpis:")
    wb.active = 0
    wb.save(path)


# ---------------------------------------------------------------- Q&A log
def qa_log(path):
    wb = Workbook()
    navod(wb, "Q&A log výběrového řízení SENTINEL", [
        "# Pravidla",
        "1) Dotazy přijímá pouze kontaktní osoba Banky e-mailem na adresu uvedenou v Zadávací dokumentaci, kap. 8.",
        "2) Každý dotaz dostane ID (Q-NNN) a zapíše se do listu Interní evidence (kdo se ptal – NIKDY se nesdílí).",
        "3) Odpověď schvaluje vlastník dotčeného požadavku (IT / Compliance / Legal / Obchod) a Procurement.",
        "4) Do listu Veřejný log se zapíše dotaz BEZ identifikace tazatele (anonymizovat i formulace, které tazatele prozrazují).",
        "5) Veřejný log se rozesílá VŠEM uchazečům současně, nejpozději 1× týdně a vždy do 2 pracovních dnů od schválení odpovědi.",
        "6) Pokud odpověď mění zadání (např. požadavek v Matici shody), vydá se nová verze dokumentu a uvede se ve sloupci „Verze dokumentu“.",
        "7) Dotazy po termínu (20. 11. 2026, 12:00) se nezodpovídají, aby nevznikla informační asymetrie.",
        "Ukázkové záznamy jsou fiktivní.",
    ])
    v = wb.create_sheet("Veřejný log")
    titul(v, "Veřejný Q&A log – sdílí se všem uchazečům")
    hlavicka(v, 5, ["ID", "Přijato", "Dokument / kapitola", "ID požadavku", "Dotaz (anonymizovaný)", "Odpověď Banky", "Zveřejněno", "Mění zadání?", "Verze dokumentu"])
    sirky(v, [8, 11, 20, 12, 56, 64, 11, 10, 14])
    qa = [
        ("Q-001", "2026-11-04", "Příloha 1 – Matice", "NFR-006", "Je přípustné řešení provozované v privátním cloudu dodavatele s datovým centrem v EU?",
         "Ano, pokud jde o dedikovanou (single-tenant) instanci, všechna data včetně záloh jsou v EU/EHP a přístup podpory z třetích zemí je vyloučen nebo předem schválen Bankou (LEG-001).", "2026-11-06", "Ne", "—"),
        ("Q-002", "2026-11-05", "Příloha 3 – TCO", "A.9", "Má být v ceně zahrnut datový feed PEP / adverse media, nebo jej Banka nakupuje samostatně?",
         "Datové zdroje musí být oceněny v listu A, řádek A.9, po celých 5 let. Pokud dodavatel podporuje více zdrojů, ocení zdroj, který doporučuje, a ostatní uvede v poznámce.", "2026-11-06", "Ne", "—"),
        ("Q-003", "2026-11-09", "Příloha 1 – Matice", "INT-002", "Jaká je verze a konfigurace bankovní Kafka platformy a jaký formát zpráv používá core banking?",
         "Apache Kafka 3.x v on-premise clusteru, zprávy ve formátu Avro se schema registry, autentizace mTLS. Popis topiců dostanou uchazeči po podpisu NDA jako Dokument D-04.", "2026-11-13", "Ne", "—"),
        ("Q-004", "2026-11-10", "Zadávací dokumentace, kap. 5", "—", "Je možné posunout termín podání nabídek o 2 týdny?",
         "Ne. Harmonogram zůstává beze změny.", "2026-11-13", "Ne", "—"),
        ("Q-005", "2026-11-12", "Příloha 1 – Matice", "TM-001", "Musí knihovna 40 scénářů pokrývat i obchodování s cennými papíry?",
         "Ano. Upřesňujeme požadavek TM-001: knihovna musí pokrývat i transakce s investičními nástroji (převody cenných papírů mezi depozitáři).", "2026-11-13", "Ano", "Matice v1.1"),
        ("Q-006", "2026-11-16", "Příloha 4 – SLA", "—", "Lze měřit dostupnost na úrovni infrastruktury dodavatele místo syntetických transakcí Banky?",
         "Ne. Dostupnost se měří end-to-end syntetickými transakcemi z monitoringu Banky (Příloha 4, kap. 3). Dodavatel může nabídnout vlastní monitoring jako doplněk.", "2026-11-20", "Ne", "—"),
        ("Q-007", "2026-11-18", "Příloha 2 – FS", "kap. 7", "Jaký objem dat se bude migrovat?",
         "Cca 52 000 subjektů (aktivní i ukončení klienti za 10 let), 410 000 alertů, 23 000 případů, 14 mil. transakcí za 10 let. Detail v Dokumentu D-06 (datový model současného systému).", "2026-11-20", "Ne", "—"),
        ("Q-008", "2026-11-19", "Příloha 5 – Smluvní", "LEG-008", "Je limit odpovědnosti 200 % ročních plateb vyjednatelný?",
         "Výhrady k limitu odpovědnosti uveďte v Matici jako „Akceptuje s výhradou“ s návrhem. Budou předmětem jednání v kole BAFO. Limit pro porušení mlčenlivosti a ochrany osobních údajů vyjednatelný není.", "2026-11-20", "Ne", "—"),
    ]
    for i, row in enumerate(qa):
        for j, val in enumerate(row):
            c = v.cell(row=6 + i, column=j + 1, value=val)
            c.alignment = WRAP
            c.border = BOX
    for i in range(len(qa), len(qa) + 30):
        for j in range(9):
            v.cell(row=6 + i, column=j + 1).border = BOX
    dv_list(v, ["Ano", "Ne"], "H6:H80")
    v.freeze_panes = "E6"
    v.conditional_formatting.add("H6:H80", CellIsRule(operator="equal", formula=['"Ano"'], fill=F_AMBER))

    n = wb.create_sheet("Interní evidence")
    titul(n, "Interní evidence dotazů – NESDÍLET (identita tazatele)")
    hlavicka(n, 5, ["ID", "Tazatel (dodavatel)", "Kontaktní osoba", "Kanál", "Vlastník odpovědi", "Schválil (Procurement)", "Stav", "Poznámka"])
    sirky(n, [8, 16, 20, 10, 18, 22, 12, 40])
    intr = [("Q-001", "Dodavatel C", "[kontakt]", "e-mail", "IT – architektura", "Ing. Pavel Černý", "Zveřejněno", ""),
            ("Q-002", "Dodavatel A", "[kontakt]", "e-mail", "IT + Procurement", "Ing. Pavel Černý", "Zveřejněno", ""),
            ("Q-003", "Dodavatel B", "[kontakt]", "e-mail", "IT – integrace", "Ing. Pavel Černý", "Zveřejněno", "Dokument D-04 předán všem po NDA"),
            ("Q-004", "Dodavatel B", "[kontakt]", "e-mail", "Procurement", "Ing. Pavel Černý", "Zveřejněno", ""),
            ("Q-005", "Dodavatel A", "[kontakt]", "e-mail", "Compliance", "Ing. Pavel Černý", "Zveřejněno", "Změna matice schválena vlastníkem a SteerCo per rollam"),
            ("Q-006", "Dodavatel C", "[kontakt]", "e-mail", "IT – provoz", "Ing. Pavel Černý", "Zveřejněno", ""),
            ("Q-007", "Dodavatel A", "[kontakt]", "e-mail", "IT – data", "Ing. Pavel Černý", "Zveřejněno", ""),
            ("Q-008", "Dodavatel C", "[kontakt]", "e-mail", "Legal", "Ing. Pavel Černý", "Zveřejněno", "")]
    for i, row in enumerate(intr):
        for j, val in enumerate(row):
            c = n.cell(row=6 + i, column=j + 1, value=val)
            c.alignment = WRAP
            c.border = BOX
            c.fill = F_RED if j == 1 else PatternFill()
    n.sheet_properties.tabColor = "C00000"
    wb.active = 1
    wb.save(path)


if __name__ == "__main__":
    matice(os.path.join(OUT, "02_Matice_shody_SABLONA.xlsx"), P.UKAZKY, "I",
           "Matice shody – ŠABLONA pro sběr požadavků (IT, Compliance, Legal, Obchod)", prazdnych=150)
    matice(os.path.join(OUT, "03_Matice_shody_VYPLNENA.xlsx"), P.POZADAVKY, "I",
           "Matice shody – fiktivně vyplněná, verze 1.0 schválená SteerCo")
    matice(os.path.join(OUT, "05_Priloha1_Matice_shody_pro_dodavatele.xlsx"),
           [r for r in P.POZADAVKY if r[9] != "Zamítnuto"], "V",
           "Příloha 1 – Matice shody (vyplňuje uchazeč)")
    tco(os.path.join(OUT, "07_Priloha3_Cenova_sablona_TCO.xlsx"))
    hodnotici_model(os.path.join(OUT, "11_Hodnotici_model_a_arch.xlsx"))
    qa_log(os.path.join(OUT, "12_QA_log.xlsx"))
    print("Hotovo:", OUT)
